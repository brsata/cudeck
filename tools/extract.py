# -*- coding: utf-8 -*-
"""Dump an existing deck to Markdown, so its content can be rebuilt.

The point of this is the speaker notes. When an old deck is redesigned, the
slides are worth rethinking and the notes are worth keeping exactly — they are
the part that took the longest to write and the part nobody can reconstruct.

    python extract.py old/*.pptx              # one .md beside each .pptx
    python extract.py old/lecture1.pptx -o notes/

Each slide becomes a heading, the text on it, and its notes verbatim. Tables
come out as Markdown tables; pictures are listed by name so you know something
was there.
"""
import argparse
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def cell_text(cell):
    return " ".join(cell.text.split())


def shape_lines(shape):
    """One shape's content, as Markdown-ish lines."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        # a hand-made diagram is usually a group, and its labels are content
        out = []
        for inner in shape.shapes:
            out += shape_lines(inner)
        return out
    if shape.has_table:
        rows = [[cell_text(c) for c in row.cells] for row in shape.table.rows]
        if not rows:
            return []
        out = ["| " + " | ".join(rows[0]) + " |",
               "|" + "|".join(["---"] * len(rows[0])) + "|"]
        out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        return out
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return ["_[picture: %s]_" % shape.name]
    if shape.has_text_frame and shape.text_frame.text.strip():
        return [line.rstrip() for line in shape.text_frame.text.splitlines()
                if line.strip()]
    return []


def slide_title(slide, body):
    """The slide's title placeholder if it has one, else its first line."""
    title = slide.shapes.title
    if title is not None and title.text_frame.text.strip():
        return " ".join(title.text_frame.text.split())
    return body[0] if body else "(no text)"


def slide_notes(slide):
    """The speaker notes, verbatim. A notes page can lack a body entirely."""
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return frame.text.strip() if frame is not None else ""


def extract(path):
    prs = Presentation(path)
    name = os.path.splitext(os.path.basename(path))[0]
    out = ["# %s" % name, "",
           "%d slides. Extracted from `%s`." % (len(prs.slides),
                                                os.path.basename(path)), ""]
    for n, slide in enumerate(prs.slides, 1):
        body = []
        for shape in slide.shapes:
            body += shape_lines(shape)
        heading = slide_title(slide, body)
        out += ["---", "", "## Slide %d — %s" % (n, heading), "",
                "**On screen**", "", "```"]
        out += body or ["(nothing)"]
        out += ["```", ""]
        out += ["**Notes**", "", slide_notes(slide) or "_(none)_", ""]
    return "\n".join(out) + "\n"


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("decks", nargs="+", help=".pptx files to read")
    ap.add_argument("-o", "--out", default=None,
                    help="directory for the .md files "
                         "(default: beside each deck)")
    args = ap.parse_args()

    for path in args.decks:
        if not path.lower().endswith(".pptx"):
            print("skipping %s" % path, file=sys.stderr)
            continue
        text = extract(path)
        folder = args.out or os.path.dirname(os.path.abspath(path))
        if args.out:
            os.makedirs(args.out, exist_ok=True)
        dest = os.path.join(
            folder, os.path.splitext(os.path.basename(path))[0] + ".md")
        with open(dest, "w", encoding="utf-8") as fh:
            fh.write(text)
        print("%s  (%d slides)" % (dest, text.count("\n## Slide ")))


if __name__ == "__main__":
    main()
