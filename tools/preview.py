# -*- coding: utf-8 -*-
"""Render a .pptx to PNG so layout can be checked without PowerPoint.

This is a proofing tool, not a renderer. It draws the shape vocabulary these
decks actually use — rectangles, flowchart nodes, text boxes, tables, pictures,
connectors — and substitutes installed fonts for the deck's own. Geometry is
exact; type is approximate, so use it to catch overflow and collisions, not to
judge letterforms.

    python preview.py deck.pptx            # all slides
    python preview.py deck.pptx 3 14 20    # just those
"""
import os
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.oxml.ns import qn

from cudeck.theme import _FONT_DIRS

DPI = 110.0
E = 914400.0                       # EMU per inch
PT = 12700.0                       # EMU per point

# the deck's own faces, once install_fonts.sh has been run, looked for where
# the theme looks for them
FONTS = {
    ("sans", False): "Archivo-Regular.ttf",
    ("sans", True): "Archivo-Bold.ttf",
    ("mono", False): "IBMPlexMono-Regular.ttf",
    ("mono", True): "IBMPlexMono-Bold.ttf",
    ("display", False): "Archivo-ExtraBold.ttf",
    ("display", True): "Archivo-ExtraBold.ttf",
}
# otherwise whatever the system has, first match wins: macOS, Linux, Windows
FALLBACK = {
    ("sans", False): ["/System/Library/Fonts/Helvetica.ttc",
                      "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                      "C:/Windows/Fonts/arial.ttf"],
    ("sans", True): ["/System/Library/Fonts/Helvetica.ttc",
                     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
                     "C:/Windows/Fonts/arialbd.ttf"],
    ("mono", False): ["/System/Library/Fonts/Menlo.ttc",
                      "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
                      "C:/Windows/Fonts/consola.ttf"],
    ("mono", True): ["/System/Library/Fonts/Menlo.ttc",
                     "/usr/share/fonts/truetype/dejavu/"
                     "DejaVuSansMono-Bold.ttf",
                     "C:/Windows/Fonts/consolab.ttf"],
}
FALLBACK[("display", False)] = FALLBACK[("display", True)] = \
    FALLBACK[("sans", True)]
TTC_INDEX = {}
_cache = {}


def _find(name):
    for d in _FONT_DIRS:
        path = os.path.join(d, name)
        if os.path.exists(path):
            return path
    return None


def font(kind, bold, pt):
    key = (kind, bold, round(pt))
    if key not in _cache:
        path = _find(FONTS[(kind, bold)])
        if path is None:
            path = next((p for p in FALLBACK[(kind, bold)]
                         if os.path.exists(p)), "")
        px = max(6, int(round(pt * DPI / 72.0)))
        idx = TTC_INDEX.get((kind, bold), 0)
        try:
            _cache[key] = ImageFont.truetype(path, px, index=idx)
        except Exception:
            _cache[key] = ImageFont.load_default()
    return _cache[key]


def px(emu):
    return emu * DPI / E


def rgb_of(el):
    c = el.find(qn("a:srgbClr"))
    if c is None:
        return None
    v = c.get("val")
    return (int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))


def shape_fill(sp):
    spPr = sp._element.find(qn("p:spPr"))
    if spPr is None:
        return None
    if spPr.find(qn("a:noFill")) is not None:
        return None
    sf = spPr.find(qn("a:solidFill"))
    return rgb_of(sf) if sf is not None else None


def shape_line(sp):
    spPr = sp._element.find(qn("p:spPr"))
    if spPr is None:
        return None, 0
    ln = spPr.find(qn("a:ln"))
    if ln is None or ln.find(qn("a:noFill")) is not None:
        return None, 0
    sf = ln.find(qn("a:solidFill"))
    if sf is None:
        return None, 0
    w = int(ln.get("w") or 12700) / PT
    return rgb_of(sf), w


def run_props(r):
    rPr = r._r.find(qn("a:rPr"))
    size, bold, color, name = 18.0, False, (0, 0, 0), ""
    if rPr is not None:
        if rPr.get("sz"):
            size = int(rPr.get("sz")) / 100.0
        bold = rPr.get("b") == "1"
        sf = rPr.find(qn("a:solidFill"))
        if sf is not None:
            color = rgb_of(sf) or color
        lt = rPr.find(qn("a:latin"))
        if lt is not None:
            name = lt.get("typeface") or ""
    kind = "mono" if ("Mono" in name or "Consolas" in name) else "sans"
    if "ExtraBold" in name:        # the headline face, much wider than Regular
        kind = "display"
    return size, bold, color, kind


# Single spacing, as a renderer draws it: 1.2x the point size for both faces,
# measured from LibreOffice. A percentage is a multiple of that, not of the
# point size, which is why the theme sets its spacing in points.
SINGLE = 1.2


def _points(el, tag):
    """A paragraph spacing given in points, or None."""
    sp = el.find(qn(tag)) if el is not None else None
    pts = sp.find(qn("a:spcPts")) if sp is not None else None
    return int(pts.get("val")) / 100.0 if pts is not None else None


def para_props(p):
    """align, space before and after (pt), line spacing, left margin.

    Line spacing is ("pts", points) or ("pct", multiple of single).
    """
    pPr = p._p.find(qn("a:pPr"))
    align, marL, line = "l", 0.0, ("pct", 1.0)
    if pPr is not None:
        align = {"ctr": "c", "r": "r"}.get(pPr.get("algn"), "l")
        marL = float(pPr.get("marL") or 0)
        ln = pPr.find(qn("a:lnSpc"))
        if ln is not None:
            pts, pct = ln.find(qn("a:spcPts")), ln.find(qn("a:spcPct"))
            if pts is not None:
                line = ("pts", int(pts.get("val")) / 100.0)
            elif pct is not None:
                line = ("pct", int(pct.get("val")) / 100000.0)
    before = _points(pPr, "a:spcBef") or 0.0
    after = _points(pPr, "a:spcAft") or 0.0
    return align, before, after, line, marL


def wrap(text, fnt, width):
    if width <= 0:
        return [text]
    out, line = [], ""
    for word in text.split(" "):
        trial = (line + " " + word) if line else word
        if fnt.getlength(trial) <= width or not line:
            line = trial
        else:
            out.append(line)
            line = word
    out.append(line)
    return out


def draw_text_frame(d, tf, box, wrap_on=True, vert=None):
    x0, y0, w, h = box
    lines = []                      # ([(text, font, colour)], align, indent)
    heights = []
    for p in tf.paragraphs:
        align, before, after, line, marL = para_props(p)
        if before:
            heights.append(("gap", before * DPI / 72.0))
        runs = [r for r in p.runs if r.text]
        if not runs:
            heights.append(("gap", after * DPI / 72.0))
            continue
        indent = px(marL)
        avail = w - indent
        toks, maxpt = [], 0
        for r in runs:
            sz, bd, col, kd = run_props(r)
            maxpt = max(maxpt, sz)
            f = font(kd, bd, sz)
            parts = r.text.split(" ")
            for i, part in enumerate(parts):
                if i:
                    toks.append((" ", f, col))
                if part:
                    toks.append((part, f, col))
        if line[0] == "pts":
            lh = line[1] * DPI / 72.0
        else:
            lh = maxpt * SINGLE * line[1] * DPI / 72.0
        if not wrap_on:
            lines.append((toks, align, indent))
            heights.append(("line", lh))
        else:
            cur, cw_ = [], 0.0
            for t, f, col in toks:
                tw = f.getlength(t)
                if cur and cw_ + tw > avail and t != " ":
                    lines.append((cur, align, indent))
                    heights.append(("line", lh))
                    cur, cw_ = [], 0.0
                    if t == " ":
                        continue
                cur.append((t, f, col))
                cw_ += tw
            if cur:
                lines.append((cur, align, indent))
                heights.append(("line", lh))
        heights.append(("gap", after * DPI / 72.0))

    total = sum(v for _, v in heights)
    y = y0
    if vert == "mid":
        y = y0 + max(0, (h - total) / 2.0)
    elif vert == "bot":
        y = y0 + max(0, h - total)

    li = 0
    for kind_, v in heights:
        if kind_ == "gap":
            y += v
            continue
        segs, align, indent = lines[li]
        li += 1
        tw = sum(f.getlength(t) for t, f, _ in segs)
        if align == "c":
            x = x0 + (w - tw) / 2.0
        elif align == "r":
            x = x0 + w - tw
        else:
            x = x0 + indent
        # the baseline sits where it would in the font's own line box, scaled
        # to the line's actual height: tight spacing pulls it up, loose pushes
        # it down
        asc, desc = max((f.getmetrics() for _, f, _ in segs),
                        key=lambda m: m[0] + m[1])
        base = y + v * asc / float(asc + desc)
        for t, f, col in segs:
            d.text((x, base), t, font=f, fill=col, anchor="ls")
            x += f.getlength(t)
        y += v


def cell_borders(cell):
    tcPr = cell._tc.find(qn("a:tcPr"))
    out = {}
    if tcPr is None:
        return out
    for edge, tag in (("B", "a:lnB"), ("T", "a:lnT")):
        ln = tcPr.find(qn(tag))
        if ln is None:
            continue
        sf = ln.find(qn("a:solidFill"))
        if sf is None:
            continue
        out[edge] = (rgb_of(sf), int(ln.get("w") or 12700) / PT)
    return out


def render(prs, idx, path):
    W = int(px(prs.slide_width))
    H = int(px(prs.slide_height))
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    slide = prs.slides[idx]

    for sp in slide.shapes:
        if sp.left is None:
            continue
        x, y = px(sp.left), px(sp.top)
        w, h = px(sp.width or 0), px(sp.height or 0)

        if sp.shape_type is not None and sp.shape_type == 13:        # PICTURE
            try:
                buf = __import__("io").BytesIO(sp.image.blob)
                im = Image.open(buf).convert("RGBA")
                im = im.resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS)
                img.paste(im, (int(x), int(y)), im)
            except Exception:
                pass
            continue

        if sp.has_table:
            t = sp.table
            cw = [px(c.width) for c in t.columns]
            rh = [px(r.height) for r in t.rows]
            ry = y
            for i, row in enumerate(t.rows):
                cx = x
                for j, cell in enumerate(row.cells):
                    for edge, (col, bw) in cell_borders(cell).items():
                        yy = ry + rh[i] if edge == "B" else ry
                        d.line([(cx, yy), (cx + cw[j], yy)], fill=col,
                               width=max(1, int(bw * DPI / 72.0)))
                    anchor = "mid"
                    tcPr = cell._tc.find(qn("a:tcPr"))
                    if tcPr is not None and tcPr.get("anchor") == "b":
                        anchor = "bot"
                    box = (cx + px(cell.margin_left), ry + 3,
                           cw[j] - px(cell.margin_left) * 2, rh[i] - 6)
                    draw_text_frame(d, cell.text_frame, box, vert=anchor)
                    cx += cw[j]
                ry += rh[i]
            continue

        fill = shape_fill(sp)
        line, lw = shape_line(sp)
        name = sp._element.find(qn("p:spPr"))
        geom = ""
        if name is not None:
            pg = name.find(qn("a:prstGeom"))
            if pg is not None:
                geom = pg.get("prst") or ""

        if geom == "flowChartDecision":
            d.polygon([(x + w / 2, y), (x + w, y + h / 2), (x + w / 2, y + h),
                       (x, y + h / 2)], fill=fill, outline=line)
        elif geom == "flowChartInputOutput":
            k = w * 0.14
            d.polygon([(x + k, y), (x + w, y), (x + w - k, y + h), (x, y + h)],
                      fill=fill, outline=line)
        elif geom in ("line", "straightConnector1"):
            d.line([(x, y), (x + w, y + h)], fill=line or (0, 0, 0),
                   width=max(1, int(lw * DPI / 72.0)))
        elif fill or line:
            d.rectangle([x, y, x + w, y + h], fill=fill, outline=line,
                        width=max(1, int(lw * DPI / 72.0)) if line else 0)

        if sp.has_text_frame and sp.text_frame.text.strip():
            tf = sp.text_frame
            bodyPr = tf._txBody.bodyPr
            rot = bodyPr.get("vert")
            anchor = {"ctr": "mid", "b": "bot"}.get(bodyPr.get("anchor"), None)
            ml, mr = px(tf.margin_left), px(tf.margin_right)
            mt = px(tf.margin_top)
            if rot == "vert270":
                box_wh = (max(1, int(h)), max(1, int(w)))
                tmp = Image.new("RGBA", box_wh, (0, 0, 0, 0))
                td = ImageDraw.Draw(tmp)
                draw_text_frame(td, tf, (0, 0, h, w), wrap_on=False, vert=None)
                tmp = tmp.rotate(90, expand=True)
                img.paste(tmp, (int(x), int(y)), tmp)
            else:
                draw_text_frame(d, tf, (x + ml, y + mt, w - ml - mr, h),
                                wrap_on=tf.word_wrap is not False, vert=anchor)

    img.save(path)
    return path


def main():
    src = sys.argv[1]
    prs = Presentation(src)
    wanted = [int(a) for a in sys.argv[2:]] or range(1, len(prs.slides) + 1)
    out = os.path.join(os.environ.get("OUT", "/tmp/cudeck"), "preview",
                       os.path.splitext(os.path.basename(src))[0])
    os.makedirs(out, exist_ok=True)
    for n in wanted:
        p = render(prs, n - 1, os.path.join(out, "s%02d.png" % n))
        print(p)


if __name__ == "__main__":
    main()
